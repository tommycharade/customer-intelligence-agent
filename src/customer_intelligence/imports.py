import csv
import hashlib
import re
import zipfile
from email import policy
from email.parser import BytesParser
from io import BytesIO, StringIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .evidence import canonical_domain
from .models import Source

MAX_UPLOAD = 10 * 1024 * 1024


def preview(filename, content):
    if len(content) > MAX_UPLOAD:
        raise ValueError("Upload a file smaller than 10 MB.")
    suffix = Path(filename).suffix.lower()
    domain = None
    published = None
    source_type = "other"
    if suffix in {".txt", ".md"}:
        text = content.decode("utf-8-sig")
    elif suffix == ".pdf":
        pdf = PdfReader(BytesIO(content))
        if pdf.is_encrypted:
            raise ValueError("Upload an unencrypted copy of this PDF.")
        if len(pdf.pages) > 100:
            raise ValueError("Split this PDF into files of at most 100 pages.")
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    elif suffix == ".docx":
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("This document expands beyond the 50 MB limit.")
        document = Document(BytesIO(content))
        text = "\n".join(
            [p.text for p in document.paragraphs]
            + [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        )
    elif suffix == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(content)
        part = message.get_body(preferencelist=("plain", "html"))
        text = str(part.get_content()) if part else ""
        if part and part.get_content_type() == "text/html":
            from .sources import PublicReader

            text = PublicReader.extract(text).get("text", "")
        text = f"Subject: {message.get('Subject', '')}\nFrom: {message.get('From', '')}\nDate: {message.get('Date', '')}\n\n{text}"
        match = re.search(r"@([a-zA-Z0-9.-]+)", str(message.get("From", "")))
        if match:
            try:
                domain = canonical_domain(match[1])
            except ValueError:
                pass
        source_type = "inbound"
    elif suffix == ".csv":
        rows = list(csv.DictReader(StringIO(content.decode("utf-8-sig"))))
        if len(rows) > 500:
            raise ValueError("Import at most 500 rows at a time.")
        sources = []
        for row in rows:
            domain_value = row.get("domain") or row.get("company_domain")
            account_domain = canonical_domain(domain_value) if domain_value else None
            sources.append(
                Source(
                    title=row.get("company") or row.get("title") or filename,
                    text="\n".join(f"{key}: {value}" for key, value in row.items()),
                    account_domain=account_domain,
                    source_type="exclusion"
                    if row.get("exclude", "").lower() in {"true", "yes", "1"}
                    else "inbound",
                )
            )
        return sources
    else:
        raise ValueError("Supported files: TXT, Markdown, PDF, DOCX, EML and CSV.")
    if not text.strip():
        raise ValueError("No readable text found. For scans or recordings, upload a text transcript.")
    if len(text) > 150000:
        raise ValueError("Split this input into smaller files (150,000 characters maximum).")
    return [
        Source(
            title=filename,
            text=text,
            source_type=source_type,
            account_domain=domain,
            published_at=published,
            content_hash=hashlib.sha256(text.encode()).hexdigest(),
        )
    ]
